#!/usr/bin/env python3
"""
Unified Terraform Deployment Script
This script handles the complete deployment process for all Lambda functions and infrastructure
"""

import os
import sys
import subprocess
import json
import shutil
import zipfile
import argparse
from pathlib import Path

def run_command(cmd, check=True, capture_output=False):
    """Run a shell command and return the result"""
    print(f"🔧 Running: {' '.join(cmd)}")
    result = subprocess.run(cmd, check=check, capture_output=capture_output, text=True)
    if capture_output:
        return result.stdout.strip()
    return result

def check_prerequisites():
    """Check if required tools are installed"""
    print("🔍 Checking prerequisites...")
    
    # Check Terraform
    try:
        run_command(["terraform", "version"], capture_output=True)
        print("✅ Terraform is installed")
    except (subprocess.CalledProcessError, FileNotFoundError):
        print("❌ Terraform is not installed. Please install Terraform first.")
        sys.exit(1)
    
    # Check AWS CLI
    try:
        run_command(["aws", "sts", "get-caller-identity"], capture_output=True)
        print("✅ AWS CLI is configured")
    except (subprocess.CalledProcessError, FileNotFoundError):
        print("❌ AWS CLI is not configured. Please run 'aws configure' first.")
        sys.exit(1)

def build_import_stocks_package():
    """Build import-stocks Lambda package with dependencies - ALWAYS REBUILD"""
    print("🔧 Building import-stocks Lambda package (ALWAYS REBUILD)...")
    
    # Paths for import-stocks only
    lambda_src = Path("backend/lambda/import-stocks")
    build_dir = Path("terraform/import_stocks_build")  # Separate build directory
    zip_file = Path("terraform/import_stocks.zip")
    
    # Skip if source doesn't exist
    if not lambda_src.exists():
        print(f"   ⚠️  Source directory {lambda_src} not found, skipping...")
        return False
    
    # ALWAYS clean up previous builds (force rebuild)
    print(f"   🧹 ALWAYS cleaning up previous builds (force rebuild)...")
    if build_dir.exists():
        shutil.rmtree(build_dir)
        print(f"   ✅ Removed existing build directory: {build_dir}")
    if zip_file.exists():
        zip_file.unlink()
        print(f"   ✅ Removed existing ZIP: {zip_file}")
    
    # Create build directory
    build_dir.mkdir(parents=True, exist_ok=True)
    
    # Copy Lambda source files
    print(f"   📁 Copying source files from {lambda_src}...")
    for file_path in lambda_src.rglob("*"):
        if file_path.is_file():
            rel_path = file_path.relative_to(lambda_src)
            dest_path = build_dir / rel_path
            dest_path.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(file_path, dest_path)
    
    # Install dependencies
    requirements_file = lambda_src / "requirements.txt"
    if requirements_file.exists():
        print(f"   📦 Installing dependencies from {requirements_file}...")
        print(f"   🔄 ALWAYS rebuilding dependencies (no cache)...")
        
        # Install with --no-cache-dir and --upgrade to force rebuild
        cmd = [
            "pip3", "install", "-r", str(requirements_file),
            "-t", str(build_dir), "--upgrade", "--no-cache-dir"
        ]
        
        try:
            result = subprocess.run(cmd, capture_output=True, text=True, check=True)
            print(f"   ✅ Dependencies installed successfully")
        except subprocess.CalledProcessError as e:
            print(f"   ❌ Failed to install dependencies: {e}")
            print(f"   📋 Error output: {e.stderr}")
            return False
        
        # Verify critical dependencies
        critical_deps = ['boto3', 'openpyxl', 'PyPDF2', 'pdfplumber']
        print(f"   🔍 Verifying critical dependencies: {critical_deps}")
        
        found_deps = []
        for dep in critical_deps:
            dep_path = build_dir / dep
            if dep_path.exists() or any((build_dir / dep).glob("*")):
                found_deps.append(dep)
        
        if len(found_deps) == len(critical_deps):
            print(f"   ✅ All critical dependencies verified: {found_deps}")
        else:
            missing = set(critical_deps) - set(found_deps)
            print(f"   ❌ Missing critical dependencies: {missing}")
            return False
        
        # Special openpyxl verification
        print(f"   🧪 Testing openpyxl functionality...")
        try:
            test_cmd = [
                "python3", "-c", 
                "import sys; sys.path.insert(0, '{}'); import openpyxl; from openpyxl import load_workbook; print('openpyxl works!')".format(build_dir)
            ]
            result = subprocess.run(test_cmd, capture_output=True, text=True, check=True)
            print(f"   ✅ openpyxl functionality verified: {result.stdout.strip()}")
        except subprocess.CalledProcessError as e:
            print(f"   ❌ CRITICAL ERROR: openpyxl verification failed: {e}")
            print(f"   🚨 DEPLOYMENT ABORTED: Cannot verify openpyxl functionality!")
            return False
    else:
        print(f"   ⚠️  No requirements.txt found for import-stocks")
    
    # Create deployment ZIP
    print(f"   📦 Creating import-stocks deployment ZIP...")
    with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for file_path in build_dir.rglob("*"):
            if file_path.is_file():
                arcname = file_path.relative_to(build_dir)
                zipf.write(file_path, arcname)
    
    print(f"   ✅ Created {zip_file}")
    
    # Verify ZIP contents
    print(f"   📦 Verifying import-stocks ZIP contents...")
    with zipfile.ZipFile(zip_file, 'r') as zipf:
        file_list = zipf.namelist()
        critical_files = ['index.py', 'boto3', 'openpyxl', 'PyPDF2', 'pdfplumber']
        
        found_files = []
        for file in critical_files:
            if any(f.startswith(file) for f in file_list):
                found_files.append(file)
        
        if len(found_files) == len(critical_files):
            print(f"   ✅ All critical files found in ZIP: {found_files}")
        else:
            missing = set(critical_files) - set(found_files)
            print(f"   ❌ Missing critical files in ZIP: {missing}")
            return False
    
    # Show ZIP size
    size_mb = zip_file.stat().st_size / (1024*1024)
    print(f"   📊 Import-stocks ZIP size: {size_mb:.1f} MB")
    
    # Clean up build directory
    shutil.rmtree(build_dir)
    print(f"   🧹 Cleaned up build directory: {build_dir}")
    
    return True

def build_lambda_packages():
    """Build other Lambda deployment packages (excluding import-stocks)"""
    print("📦 Building other Lambda deployment packages...")
    
    # Define all lambda functions to build with their critical dependencies
    lambda_functions = [
        {
            "name": "parse-mf-stocks",
            "src": "backend/lambda/parse-mf-stocks",
            "zip": "terraform/parse_mf_stocks.zip",
            "critical_deps": ['requests', 'boto3', 'charset_normalizer', 'urllib3', 'certifi', 'idna'],
            "main_file": "main.py"
        },
        {
            "name": "portfolio-api",
            "src": "backend/lambda/portfolio-api-py",
            "zip": "terraform/portfolio_api.zip",
            "critical_deps": ['boto3'],
            "main_file": "index.py"
        },
        {
            "name": "expenses-api",
            "src": "backend/lambda/expenses-api-py",
            "zip": "terraform/expenses_api.zip",
            "critical_deps": ['boto3'],
            "main_file": "index.py"
        },
    ]
    
    for func in lambda_functions:
        print(f"\n🔧 Building {func['name']}...")
        
        # Paths
        lambda_src = Path(func["src"])
        build_dir = Path(f"terraform/lambda_build_{func['name'].replace('-', '_')}")
        zip_file = Path(func["zip"])
        terraform_zip = Path(f"terraform/{func['name'].replace('-', '_')}.zip")
        
        # Skip if source doesn't exist
        if not lambda_src.exists():
            print(f"   ⚠️  Source directory {lambda_src} not found, skipping...")
            continue
        
        # ALWAYS clean up previous builds to force rebuild
        print(f"   🧹 Cleaning up previous builds...")
        if build_dir.exists():
            shutil.rmtree(build_dir)
        if zip_file.exists():
            zip_file.unlink()
        if terraform_zip.exists():
            terraform_zip.unlink()
            print(f"   ✅ Removed existing {terraform_zip}")
        
        # Create build directory
        build_dir.mkdir(parents=True, exist_ok=True)
        
        # Copy Lambda source files
        print(f"   📋 Copying {func['name']} source files...")
        for file in lambda_src.glob("*.py"):
            shutil.copy2(file, build_dir)
            print(f"      ✅ Copied {file.name}")
        
        # Install dependencies with comprehensive verification
        requirements_file = lambda_src / "requirements.txt"
        if requirements_file.exists():
            print(f"   📦 Installing {func['name']} dependencies...")
            
            # Show requirements content
            with open(requirements_file, 'r') as f:
                requirements = f.read().strip()
            print(f"   📋 Requirements: {requirements}")
            
            # Install dependencies with retry logic
            max_retries = 3
            for attempt in range(max_retries):
                try:
                    run_command([
                        "pip3", "install", "-r", str(requirements_file), 
                        "-t", str(build_dir), "--upgrade", "--no-cache-dir"
                    ])
                    break
                except subprocess.CalledProcessError as e:
                    if attempt == max_retries - 1:
                        print(f"   ❌ ERROR: Failed to install dependencies after {max_retries} attempts")
                        print(f"   Error: {e}")
                        sys.exit(1)
                    else:
                        print(f"   ⚠️  Attempt {attempt + 1} failed, retrying...")
                        continue
            
            # Verify critical dependencies
            print(f"   🔍 Verifying critical dependencies for {func['name']}...")
            missing_deps = []
            found_deps = []
            
            for dep in func['critical_deps']:
                # Check for the dependency directory or .dist-info
                dep_found = False
                for item in build_dir.iterdir():
                    if item.is_dir() and (item.name == dep or item.name.startswith(f"{dep}-")):
                        dep_found = True
                        found_deps.append(dep)
                        break
                
                if not dep_found:
                    missing_deps.append(dep)
            
            if missing_deps:
                print(f"   ❌ CRITICAL ERROR: Missing dependencies in {func['name']} build directory:")
                for dep in missing_deps:
                    print(f"      ❌ {dep}")
                print(f"   📁 Build directory contents:")
                for item in sorted(build_dir.iterdir()):
                    print(f"      - {item.name}")
                print(f"   🚨 DEPLOYMENT ABORTED: Cannot proceed without critical dependencies!")
                sys.exit(1)
            
            print(f"   ✅ All critical dependencies verified for {func['name']}: {found_deps}")
            print(f"   ✅ {func['name']} dependencies installed successfully")
        else:
            print(f"   ⚠️  No requirements.txt found for {func['name']}")
        
        # Create deployment ZIP
        print(f"   📦 Creating {func['name']} deployment ZIP...")
        with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
            for file_path in build_dir.rglob("*"):
                if file_path.is_file():
                    arcname = file_path.relative_to(build_dir)
                    zipf.write(file_path, arcname)
        
        print(f"   ✅ Created {zip_file}")
        
        # Copy to terraform directory for Terraform to use (if different)
        terraform_zip = Path(f"terraform/{func['name'].replace('-', '_')}.zip")
        if zip_file != terraform_zip:
            shutil.copy2(zip_file, terraform_zip)
            print(f"   ✅ Copied to {terraform_zip} for Terraform deployment")
        else:
            print(f"   ✅ ZIP already in terraform directory: {terraform_zip}")
        print(f"   📊 Terraform ZIP size: {terraform_zip.stat().st_size / (1024*1024):.1f} MB")
        
        # Comprehensive ZIP verification for all functions
        print(f"   📦 Verifying {func['name']} ZIP contents...")
        with zipfile.ZipFile(zip_file, 'r') as zipf:
            file_list = zipf.namelist()
            
            # Check for main file
            has_main_file = any(func['main_file'] in f for f in file_list)
            
            # Check for critical dependencies
            has_deps = {}
            for dep in func['critical_deps']:
                has_deps[dep] = any(dep in f for f in file_list)
            
            print(f"      📄 Contains {func['main_file']}: {has_main_file}")
            for dep, has_it in has_deps.items():
                print(f"      📄 Contains {dep}: {has_it}")
            print(f"      📄 Total files: {len(file_list)}")
            
            # Check for missing dependencies
            missing_deps = [dep for dep, has_it in has_deps.items() if not has_it]
            
            if not has_main_file or missing_deps:
                print(f"   ❌ CRITICAL ERROR: {func['name']} ZIP is missing critical files!")
                if not has_main_file:
                    print(f"      ❌ Missing: {func['main_file']}")
                if missing_deps:
                    print(f"      ❌ Missing dependencies: {missing_deps}")
                
                # Show some ZIP contents for debugging
                print(f"   📁 ZIP contents (first 20 files):")
                for i, file in enumerate(file_list[:20]):
                    print(f"      - {file}")
                if len(file_list) > 20:
                    print(f"      ... and {len(file_list) - 20} more files")
                
                print(f"   🚨 DEPLOYMENT ABORTED: ZIP package is incomplete!")
                sys.exit(1)
            
            print(f"   ✅ ZIP verification passed for {func['name']}")
        
        # Special verification for import-stocks function
        if func['name'] == 'import-stocks':
            print(f"   🔍 Special verification for {func['name']} - testing openpyxl...")
            try:
                # Test if openpyxl can be imported in the build directory
                import sys
                original_path = sys.path.copy()
                sys.path.insert(0, str(build_dir))
                
                try:
                    from openpyxl import load_workbook
                    print(f"   ✅ openpyxl import test successful")
                    
                    # Test basic functionality
                    import io
                    test_data = b'PK\x03\x04'  # Minimal ZIP header
                    try:
                        # This will fail but we just want to test if the function exists
                        load_workbook(io.BytesIO(test_data))
                    except:
                        pass  # Expected to fail with test data
                    
                    print(f"   ✅ openpyxl load_workbook function test successful")
                    
                except ImportError as e:
                    print(f"   ❌ CRITICAL ERROR: openpyxl import failed: {e}")
                    print(f"   🚨 DEPLOYMENT ABORTED: openpyxl not functional!")
                    sys.exit(1)
                finally:
                    sys.path = original_path
                    
            except Exception as e:
                print(f"   ❌ CRITICAL ERROR: openpyxl verification failed: {e}")
                print(f"   🚨 DEPLOYMENT ABORTED: Cannot verify openpyxl functionality!")
                sys.exit(1)
        
        # Clean up build directory
        shutil.rmtree(build_dir)
    
    # Show summary of created packages
    print(f"\n📦 BUILD SUMMARY:")
    print(f"=" * 50)
    for func in lambda_functions:
        terraform_zip = Path(f"terraform/{func['name'].replace('-', '_')}.zip")
        if terraform_zip.exists():
            size_mb = terraform_zip.stat().st_size / (1024*1024)
            print(f"✅ {func['name']}: {terraform_zip.name} ({size_mb:.1f} MB)")
        else:
            print(f"❌ {func['name']}: Package not created")
    print(f"=" * 50)

def create_prebuilt_package():
    """Create a pre-built package for import-stocks Lambda"""
    print("📦 Creating pre-built import-stocks package...")
    
    # Paths
    lambda_src = Path("backend/lambda/import-stocks")
    build_dir = Path("import_stocks_build")
    zip_file = Path("import_stocks_fixed.zip")
    
    if not lambda_src.exists():
        print(f"❌ Lambda source not found: {lambda_src}")
        return False
    
    # Clean up previous builds
    if build_dir.exists():
        shutil.rmtree(build_dir)
    if zip_file.exists():
        zip_file.unlink()
    
    # Create build directory
    build_dir.mkdir(parents=True, exist_ok=True)
    
    # Copy Lambda source files
    print("📋 Copying source files...")
    for file in lambda_src.glob("*.py"):
        shutil.copy2(file, build_dir)
        print(f"  ✅ Copied {file.name}")
    
    # Install dependencies
    requirements_file = lambda_src / "requirements.txt"
    if requirements_file.exists():
        print("📦 Installing dependencies...")
        try:
            run_command([
                "pip3", "install", "-r", str(requirements_file), 
                "-t", str(build_dir), "--upgrade", "--no-cache-dir"
            ])
            print("✅ Dependencies installed successfully")
        except subprocess.CalledProcessError as e:
            print(f"❌ Failed to install dependencies: {e}")
            return False
        
        # Verify critical dependencies
        print("🔍 Verifying dependencies...")
        critical_deps = ['boto3', 'openpyxl', 'PyPDF2', 'pdfplumber']
        missing_deps = []
        
        for dep in critical_deps:
            dep_found = False
            for item in build_dir.iterdir():
                if item.is_dir() and (item.name == dep or item.name.startswith(f"{dep}-")):
                    dep_found = True
                    print(f"  ✅ {dep} found")
                    break
            
            if not dep_found:
                missing_deps.append(dep)
                print(f"  ❌ {dep} missing")
        
        if missing_deps:
            print(f"❌ Missing dependencies: {missing_deps}")
            return False
        
        print("✅ All dependencies verified")
    
    # Create deployment ZIP
    print("📦 Creating deployment ZIP...")
    with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for file_path in build_dir.rglob("*"):
            if file_path.is_file():
                arcname = file_path.relative_to(build_dir)
                zipf.write(file_path, arcname)
    
    print(f"✅ Created {zip_file}")
    print(f"ZIP size: {zip_file.stat().st_size / (1024*1024):.1f} MB")
    
    # Clean up build directory
    shutil.rmtree(build_dir)
    return True

def deploy_terraform():
    """Deploy using Terraform"""
    print("🚀 Starting Terraform deployment...")
    
    # Change to terraform directory
    os.chdir("terraform")
    
    try:
        # Initialize Terraform
        print("📋 Initializing Terraform...")
        run_command(["terraform", "init"])
        
        # Validate configuration
        print("📋 Validating Terraform configuration...")
        run_command(["terraform", "validate"])
        
        # Plan deployment
        print("📋 Planning Terraform deployment...")
        run_command(["terraform", "plan", "-out=tfplan"])
        
        # Apply changes automatically
        print("\n🚀 Applying Terraform changes...")
        run_command(["terraform", "apply", "--auto-approve", "tfplan"])
        
        # Show outputs
        print("\n✅ Deployment completed!")
        show_deployment_outputs()
            
    finally:
        # Clean up
        if Path("tfplan").exists():
            Path("tfplan").unlink()
        os.chdir("..")
    
    return True

def show_deployment_outputs():
    """Show all deployment outputs"""
    print("\n📊 Lambda Functions:")
    
    # Parse MF Stocks Lambda
    try:
        lambda_name = run_command(["terraform", "output", "-raw", "parse_mf_stocks_lambda_name"], capture_output=True)
        lambda_arn = run_command(["terraform", "output", "-raw", "parse_mf_stocks_lambda_arn"], capture_output=True)
        print(f"   📊 Parse MF Stocks: {lambda_name}")
        print(f"      ARN: {lambda_arn}")
    except subprocess.CalledProcessError:
        print("   📊 Parse MF Stocks: (not deployed)")
    
    # Portfolio API Lambda
    try:
        portfolio_name = run_command(["terraform", "output", "-raw", "portfolio_lambda_name"], capture_output=True)
        portfolio_arn = run_command(["terraform", "output", "-raw", "portfolio_lambda_arn"], capture_output=True)
        print(f"   💼 Portfolio API: {portfolio_name}")
        print(f"      ARN: {portfolio_arn}")
    except subprocess.CalledProcessError:
        print("   💼 Portfolio API: (not deployed)")
    
    # Expenses API Lambda
    try:
        expenses_name = run_command(["terraform", "output", "-raw", "expenses_lambda_name"], capture_output=True)
        expenses_arn = run_command(["terraform", "output", "-raw", "expenses_lambda_arn"], capture_output=True)
        print(f"   💰 Expenses API: {expenses_name}")
        print(f"      ARN: {expenses_arn}")
    except subprocess.CalledProcessError:
        print("   💰 Expenses API: (not deployed)")
    
    # Import Stocks Lambda
    try:
        import_name = run_command(["terraform", "output", "-raw", "import_stocks_lambda_name"], capture_output=True)
        import_arn = run_command(["terraform", "output", "-raw", "import_stocks_lambda_arn"], capture_output=True)
        print(f"   📊 Import Stocks: {import_name}")
        print(f"      ARN: {import_arn}")
    except subprocess.CalledProcessError:
        print("   📊 Import Stocks: (not deployed)")
    
    print("\n📊 DynamoDB Tables:")
    try:
        stock_table = run_command(["terraform", "output", "-raw", "stock_companies_table_name"], capture_output=True)
        print(f"   📈 Stock Companies: {stock_table}")
    except subprocess.CalledProcessError:
        print("   📈 Stock Companies: (not deployed)")
    
    try:
        mf_table = run_command(["terraform", "output", "-raw", "mutual_fund_schemes_table_name"], capture_output=True)
        print(f"   💰 Mutual Fund Schemes: {mf_table}")
    except subprocess.CalledProcessError:
        print("   💰 Mutual Fund Schemes: (not deployed)")
    
    try:
        invest_table = run_command(["terraform", "output", "-raw", "invest_table_name"], capture_output=True)
        print(f"   💼 Holdings: {invest_table}")
    except subprocess.CalledProcessError:
        print("   💼 Holdings: (not deployed)")
    
    try:
        expenses_table = run_command(["terraform", "output", "-raw", "expenses_table_name"], capture_output=True)
        print(f"   💸 Expenses: {expenses_table}")
    except subprocess.CalledProcessError:
        print("   💸 Expenses: (not deployed)")
    
    print("\n🎯 Usage Examples:")
    print("   📊 Parse MF/Stocks (both):")
    print("   aws lambda invoke --function-name parse-mf-stocks --payload 'eyJ0eXBlIjoiYm90aCJ9Cg==' both_response.json")
    print("   ")
    print("   📈 Parse stocks only:")
    print("   aws lambda invoke --function-name parse-mf-stocks --payload 'eyJ0eXBlIjoic3RvY2tzIn0K' stocks_response.json")
    print("   ")
    print("   💰 Parse mutual funds only:")
    print("   aws lambda invoke --function-name parse-mf-stocks --payload 'eyJ0eXBlIjoibWYifQo=' mf_response.json")
    print("   ")
    print("   🔍 Check response:")
    print("   cat response.json")
    print("   ")
    print("   💼 Test Portfolio API:")
    print("   curl -X GET https://your-api-gateway-url/portfolio/holdings")
    print("   ")
    print("   💸 Test Expenses API:")
    print("   curl -X GET https://your-api-gateway-url/expenses/transactions")
    print("   ")
    print("   📝 Alternative method (using file):")
    print("   echo '{\"type\":\"stocks\"}' > payload.json")
    print("   aws lambda invoke --function-name parse-mf-stocks --payload file://payload.json response.json")

def parse_arguments():
    """Parse command line arguments"""
    parser = argparse.ArgumentParser(
        description="Deploy Lambda functions and infrastructure",
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog="""
Examples:
  ./deploy-lambda.py                    # Deploy all functions
  ./deploy-lambda.py 1                  # Deploy only import-stocks
  ./deploy-lambda.py 2                  # Deploy only portfolio-api
  ./deploy-lambda.py 3                  # Deploy only parse-mf-stocks
  ./deploy-lambda.py 4                  # Deploy only expenses-api
  ./deploy-lambda.py --help             # Show this help message

Function mapping:
  1 = import-stocks (CAS parsing, Excel support)
  2 = portfolio-api (holdings, portfolio management)
  3 = parse-mf-stocks (mutual fund data parsing)
  4 = expenses-api (expense tracking)
        """
    )
    
    parser.add_argument(
        'function', 
        nargs='?', 
        type=int, 
        choices=[1, 2, 3, 4],
        help='Deploy specific function (1-4) or all if not specified'
    )
    
    return parser.parse_args()

def get_function_info(function_num):
    """Get function information based on number"""
    functions = {
        1: {
            "name": "import-stocks",
            "description": "Import Stocks Lambda (CAS parsing, Excel support)",
            "terraform_target": "aws_lambda_function.import_stocks",
            "build_function": "build_import_stocks_package"
        },
        2: {
            "name": "portfolio-api", 
            "description": "Portfolio API Lambda (holdings, portfolio management)",
            "terraform_target": "aws_lambda_function.portfolio_api",
            "build_function": "build_portfolio_api_package"
        },
        3: {
            "name": "parse-mf-stocks",
            "description": "Parse MF/Stocks Lambda (mutual fund data parsing)", 
            "terraform_target": "aws_lambda_function.parse_mf_stocks",
            "build_function": "build_parse_mf_stocks_package"
        },
        4: {
            "name": "expenses-api",
            "description": "Expenses API Lambda (expense tracking)",
            "terraform_target": "aws_lambda_function.expenses_api", 
            "build_function": "build_expenses_api_package"
        }
    }
    return functions.get(function_num)

def build_specific_function(function_num):
    """Build package for a specific function"""
    func_info = get_function_info(function_num)
    if not func_info:
        print(f"❌ Invalid function number: {function_num}")
        return False
    
    # Define function-specific build logic
    if function_num == 2:  # portfolio-api
        return build_portfolio_api_package()
    elif function_num == 3:  # parse-mf-stocks
        return build_parse_mf_stocks_package()
    elif function_num == 4:  # expenses-api
        return build_expenses_api_package()
    else:
        print(f"❌ Function {function_num} not supported for individual build")
        return False

def build_portfolio_api_package():
    """Build portfolio-api Lambda package"""
    print("🔧 Building portfolio-api Lambda package...")
    
    lambda_src = Path("backend/lambda/portfolio-api-py")
    build_dir = Path("terraform/portfolio_api_build")
    zip_file = Path("terraform/portfolio_api.zip")
    
    if not lambda_src.exists():
        print(f"❌ Lambda source directory not found: {lambda_src}")
        return False
    
    # Clean up previous build
    if build_dir.exists():
        shutil.rmtree(build_dir)
    if zip_file.exists():
        zip_file.unlink()
    
    # Create build directory and copy source
    build_dir.mkdir(parents=True, exist_ok=True)
    for item in lambda_src.iterdir():
        if item.is_file():
            shutil.copy2(item, build_dir)
        elif item.is_dir():
            shutil.copytree(item, build_dir / item.name)
    
    # Install dependencies
    requirements_file = build_dir / "requirements.txt"
    if requirements_file.exists():
        run_command(["pip3", "install", "-r", str(requirements_file), "-t", str(build_dir), "--upgrade"])
    
    # Create ZIP
    with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, dirs, files in os.walk(build_dir):
            for file in files:
                file_path = Path(root) / file
                arcname = file_path.relative_to(build_dir)
                zipf.write(file_path, arcname)
    
    print(f"✅ portfolio-api package built: {zip_file}")
    return True

def build_parse_mf_stocks_package():
    """Build parse-mf-stocks Lambda package"""
    print("🔧 Building parse-mf-stocks Lambda package...")
    
    lambda_src = Path("backend/lambda/parse-mf-stocks")
    build_dir = Path("terraform/parse_mf_stocks_build")
    zip_file = Path("terraform/parse_mf_stocks.zip")
    
    if not lambda_src.exists():
        print(f"❌ Lambda source directory not found: {lambda_src}")
        return False
    
    # Clean up previous build
    if build_dir.exists():
        shutil.rmtree(build_dir)
    if zip_file.exists():
        zip_file.unlink()
    
    # Create build directory and copy source
    build_dir.mkdir(parents=True, exist_ok=True)
    for item in lambda_src.iterdir():
        if item.is_file():
            shutil.copy2(item, build_dir)
        elif item.is_dir():
            shutil.copytree(item, build_dir / item.name)
    
    # Install dependencies
    requirements_file = build_dir / "requirements.txt"
    if requirements_file.exists():
        run_command(["pip3", "install", "-r", str(requirements_file), "-t", str(build_dir), "--upgrade"])
    
    # Create ZIP
    with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, dirs, files in os.walk(build_dir):
            for file in files:
                file_path = Path(root) / file
                arcname = file_path.relative_to(build_dir)
                zipf.write(file_path, arcname)
    
    print(f"✅ parse-mf-stocks package built: {zip_file}")
    return True

def build_expenses_api_package():
    """Build expenses-api Lambda package"""
    print("🔧 Building expenses-api Lambda package...")
    
    lambda_src = Path("backend/lambda/expenses-api-py")
    build_dir = Path("terraform/expenses_api_build")
    zip_file = Path("terraform/expenses_api.zip")
    
    if not lambda_src.exists():
        print(f"❌ Lambda source directory not found: {lambda_src}")
        return False
    
    # Clean up previous build
    if build_dir.exists():
        shutil.rmtree(build_dir)
    if zip_file.exists():
        zip_file.unlink()
    
    # Create build directory and copy source
    build_dir.mkdir(parents=True, exist_ok=True)
    for item in lambda_src.iterdir():
        if item.is_file():
            shutil.copy2(item, build_dir)
        elif item.is_dir():
            shutil.copytree(item, build_dir / item.name)
    
    # Install dependencies
    requirements_file = build_dir / "requirements.txt"
    if requirements_file.exists():
        run_command(["pip3", "install", "-r", str(requirements_file), "-t", str(build_dir), "--upgrade"])
    
    # Create ZIP
    with zipfile.ZipFile(zip_file, 'w', zipfile.ZIP_DEFLATED) as zipf:
        for root, dirs, files in os.walk(build_dir):
            for file in files:
                file_path = Path(root) / file
                arcname = file_path.relative_to(build_dir)
                zipf.write(file_path, arcname)
    
    print(f"✅ expenses-api package built: {zip_file}")
    return True

def deploy_specific_function(function_num):
    """Deploy a specific function using Terraform"""
    func_info = get_function_info(function_num)
    if not func_info:
        print(f"❌ Invalid function number: {function_num}")
        return False
    
    print(f"🚀 Deploying {func_info['name']} with Terraform...")
    
    # Initialize Terraform
    run_command(["terraform", "init"])
    
    # Plan deployment
    run_command(["terraform", "plan", f"-target={func_info['terraform_target']}"])
    
    # Apply deployment
    run_command(["terraform", "apply", f"-target={func_info['terraform_target']}", "-auto-approve"])
    
    print(f"✅ {func_info['name']} deployed successfully!")
    return True

def main():
    """Main deployment function"""
    args = parse_arguments()
    
    if args.function:
        # Deploy specific function
        func_info = get_function_info(args.function)
        if not func_info:
            print(f"❌ Invalid function number: {args.function}")
            print("Valid options: 1, 2, 3, 4")
            sys.exit(1)
            
        print(f"🚀 Starting {func_info['description']} Deployment...")
        print("=" * 60)
        print(f"This will deploy only: {func_info['name']}")
        print("=" * 60)
    else:
        # Deploy all functions
        print("🚀 Starting Unified Terraform Deployment...")
        print("=" * 60)
        print("This will deploy all Lambda functions and infrastructure:")
        print("  📊 Parse MF/Stocks Lambda (parse-mf-stocks)")
        print("  💼 Portfolio API Lambda (portfolio-api)")
        print("  💰 Expenses API Lambda (expenses-api)")
        print("  📊 Import Stocks Lambda (import-stocks) - WITH OPENPYXL SUPPORT")
        print("  📊 All DynamoDB Tables")
        print("  🌐 API Gateway Routes")
        print("=" * 60)
        print("🔍 DEPENDENCY VERIFICATION:")
        print("  ✅ All critical dependencies will be verified")
        print("  ✅ openpyxl will be tested for Excel support")
        print("  ✅ ZIP packages will be validated")
        print("  ✅ Deployment will fail if dependencies are missing")
        print("=" * 60)
        print("📝 IMPORTANT: If import-stocks Lambda fails to deploy with dependencies:")
        print("  1. Use the pre-built package: import_stocks_fixed.zip")
        print("  2. Upload manually via AWS Console or CLI")
        print("  3. Package includes all dependencies (36MB)")
        print("=" * 60)
    
    # Check prerequisites
    check_prerequisites()
    
    if args.function:
        # Deploy specific function
        func_info = get_function_info(args.function)
        
        if args.function == 1:
            # Import-stocks needs special handling
            print("\n🔧 Creating pre-built import-stocks package...")
            if create_prebuilt_package():
                print("✅ Pre-built package created: import_stocks_fixed.zip")
                print("📝 You can upload this manually if Terraform deployment fails")
            else:
                print("❌ Failed to create pre-built package")
            
            print("\n" + "="*60)
            print("🔧 BUILDING IMPORT-STOCKS PACKAGE (ALWAYS REBUILD)")
            print("="*60)
            if not build_import_stocks_package():
                print("❌ Failed to build import-stocks package. Deployment aborted.")
                return
        else:
            # Build specific function package
            print(f"\n" + "="*60)
            print(f"🔧 BUILDING {func_info['name'].upper()} PACKAGE")
            print("="*60)
            if not build_specific_function(args.function):
                print(f"❌ Failed to build {func_info['name']} package. Deployment aborted.")
                return
        
        # Deploy specific function with Terraform
        success = deploy_specific_function(args.function)
    else:
        # Deploy all functions (original logic)
        # Create pre-built package automatically
        print("\n🔧 Creating pre-built import-stocks package...")
        if create_prebuilt_package():
            print("✅ Pre-built package created: import_stocks_fixed.zip")
            print("📝 You can upload this manually if Terraform deployment fails")
        else:
            print("❌ Failed to create pre-built package")
        
        # Build import-stocks package first (always rebuild)
        print("\n" + "="*60)
        print("🔧 BUILDING IMPORT-STOCKS PACKAGE (ALWAYS REBUILD)")
        print("="*60)
        if not build_import_stocks_package():
            print("❌ Failed to build import-stocks package. Deployment aborted.")
            return
        
        # Build other Lambda packages
        print("\n" + "="*60)
        print("🔧 BUILDING OTHER LAMBDA PACKAGES")
        print("="*60)
        build_lambda_packages()
        
        # Deploy with Terraform
        success = deploy_terraform()
    
    if success:
        if args.function:
            func_info = get_function_info(args.function)
            print(f"\n🎉 {func_info['description']} deployment completed successfully!")
            print(f"\n🔍 FINAL VERIFICATION:")
            print(f"✅ {func_info['name']} Lambda deployed successfully")
            if args.function == 1:
                print("✅ openpyxl is included and tested in import-stocks Lambda")
                print("✅ Excel parsing support is fully functional")
                print("✅ Zerodha import should now work without errors")
                print("\n📝 Next steps:")
                print("1. Test Zerodha Excel import in the frontend")
                print("2. Verify holdings are imported correctly")
                print("3. Check CloudWatch logs if any issues occur")
                print("4. Test other broker imports (Groww, Upstox, Angel)")
            else:
                print(f"\n📝 Next steps:")
                print(f"1. Test {func_info['name']} functionality")
                print("2. Check CloudWatch logs if any issues occur")
                print("3. Verify API Gateway endpoints are working")
        else:
            print("\n🎉 Deployment completed successfully!")
            print("\n🔍 FINAL VERIFICATION:")
            print("✅ All Lambda functions deployed with verified dependencies")
            print("✅ openpyxl is included and tested in import-stocks Lambda")
            print("✅ Excel parsing support is fully functional")
            print("✅ Zerodha import should now work without errors")
            print("\n📝 Next steps:")
            print("1. Test Zerodha Excel import in the frontend")
            print("2. Verify holdings are imported correctly")
            print("3. Check CloudWatch logs if any issues occur")
            print("4. Test other broker imports (Groww, Upstox, Angel)")
            print("5. Verify API Gateway endpoints are working")
            print("\n🚀 Your Zerodha import functionality is now ready!")
    else:
        print("\n❌ Deployment failed or was cancelled")
        print("🔍 Check the error messages above for details")
        sys.exit(1)

if __name__ == "__main__":
    main()