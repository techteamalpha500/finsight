#!/usr/bin/env python3
"""
Quick test to see what happens when importing openpyxl
"""

def test_openpyxl_import():
    try:
        print("Attempting to import openpyxl...")
        from openpyxl import load_workbook
        print("✅ SUCCESS: openpyxl imported successfully!")
        print(f"load_workbook function: {load_workbook}")
        return True
    except ImportError as e:
        print(f"❌ FAILED: ImportError when importing openpyxl: {e}")
        return False
    except Exception as e:
        print(f"❌ FAILED: Other error when importing openpyxl: {e}")
        return False

def handler(event, context):
    """Lambda handler for testing"""
    result = test_openpyxl_import()
    
    return {
        'statusCode': 200,
        'body': {
            'success': result,
            'message': 'openpyxl import test completed'
        }
    }

if __name__ == "__main__":
    test_openpyxl_import()
