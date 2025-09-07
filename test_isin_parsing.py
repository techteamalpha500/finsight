#!/usr/bin/env python3
"""
Test the updated ISIN parsing logic
"""

import openpyxl
from pathlib import Path

def test_isin_parsing():
    """Test the updated ISIN parsing logic"""
    excel_file = Path("holdings-GYS673.xlsx")
    
    if not excel_file.exists():
        print(f"❌ File not found: {excel_file}")
        return
    
    print(f"🧪 Testing ISIN Parsing Logic")
    print("=" * 60)
    
    try:
        # Load the workbook
        workbook = openpyxl.load_workbook(excel_file, data_only=True)
        
        # Test with Equity sheet
        if 'Equity' in workbook.sheetnames:
            sheet = workbook['Equity']
            print(f"📄 Testing with Equity Sheet")
            
            # Convert sheet to rows (simulating the lambda input)
            rows = []
            for row in sheet.iter_rows(values_only=True):
                rows.append([str(cell) if cell is not None else "" for cell in row])
            
            # Find the header row
            header_row_index = None
            for i, row in enumerate(rows):
                row_text = ' '.join([str(cell).lower() for cell in row if cell])
                if 'symbol' in row_text and 'isin' in row_text and 'sector' in row_text:
                    header_row_index = i
                    break
            
            if header_row_index is None:
                print("❌ Header row not found")
                return
            
            print(f"✅ Found header row at index: {header_row_index}")
            
            # Get header row to understand column positions
            header_row = rows[header_row_index]
            print(f"📋 Headers: {header_row}")
            
            # Find column indices
            symbol_col = None
            isin_col = None
            sector_col = None
            quantity_col = None
            avg_price_col = None
            prev_close_col = None
            
            for i, cell in enumerate(header_row):
                cell_text = str(cell).lower().strip()
                if 'symbol' in cell_text:
                    symbol_col = i
                elif 'isin' in cell_text:
                    isin_col = i
                elif 'sector' in cell_text:
                    sector_col = i
                elif 'quantity available' in cell_text:
                    quantity_col = i
                elif 'average price' in cell_text:
                    avg_price_col = i
                elif 'previous closing price' in cell_text:
                    prev_close_col = i
            
            print(f"📊 Column positions:")
            print(f"   Symbol: {symbol_col}")
            print(f"   ISIN: {isin_col}")
            print(f"   Sector: {sector_col}")
            print(f"   Quantity: {quantity_col}")
            print(f"   Average Price: {avg_price_col}")
            print(f"   Previous Close: {prev_close_col}")
            
            # Parse stock data rows
            stocks = []
            for row in rows[header_row_index + 1:]:
                # Skip empty rows
                if not any(str(cell).strip() for cell in row if cell):
                    continue
                
                # Skip summary rows
                row_text = ' '.join([str(cell).lower() for cell in row if cell])
                if any(keyword in row_text for keyword in ['invested value', 'present value', 'unrealized p&l', 'summary', 'statement']):
                    continue
                
                try:
                    # Extract stock data
                    symbol = str(row[symbol_col]).strip().upper() if symbol_col is not None and symbol_col < len(row) and row[symbol_col] else None
                    isin = str(row[isin_col]).strip().upper() if isin_col is not None and isin_col < len(row) and row[isin_col] else None
                    sector = str(row[sector_col]).strip() if sector_col is not None and sector_col < len(row) and row[sector_col] else None
                    quantity = float(str(row[quantity_col]).replace(',', '')) if quantity_col is not None and quantity_col < len(row) and row[quantity_col] else 0
                    avg_price = float(str(row[avg_price_col]).replace(',', '')) if avg_price_col is not None and avg_price_col < len(row) and row[avg_price_col] else 0
                    prev_close = float(str(row[prev_close_col]).replace(',', '')) if prev_close_col is not None and prev_close_col < len(row) and row[prev_close_col] else 0
                    
                    # Skip if no symbol, ISIN, or quantity
                    if not symbol or not isin or quantity <= 0:
                        continue
                    
                    # Use previous closing price as current price, fallback to average price
                    current_price = prev_close if prev_close > 0 else avg_price
                    current_value = quantity * current_price
                    invested_value = quantity * avg_price
                    
                    stock = {
                        'name': symbol,
                        'symbol': symbol,
                        'isin': isin,
                        'sector': sector,
                        'units': quantity,
                        'price': current_price,
                        'currentValue': current_value,
                        'investedAmount': invested_value
                    }
                    stocks.append(stock)
                    
                except (ValueError, IndexError, TypeError) as e:
                    continue
            
            print(f"\n📈 Parsed {len(stocks)} stocks with ISIN:")
            for i, stock in enumerate(stocks[:5]):  # Show first 5
                print(f"   {i+1}. {stock['symbol']} ({stock['isin']}): {stock['units']} units @ ₹{stock['price']:.2f} = ₹{stock['currentValue']:.2f}")
            
            if len(stocks) > 5:
                print(f"   ... and {len(stocks) - 5} more stocks")
            
            print(f"\n✅ SUCCESS: Parsed {len(stocks)} stocks with ISIN correctly!")
            
        else:
            print("❌ Equity sheet not found")
    
    except Exception as e:
        print(f"❌ Error: {e}")

if __name__ == "__main__":
    test_isin_parsing()